"""写 Excel：主表按模板原样填，问题表单独一个文件。"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

FIRST_DATA_ROW = 3  # 第1行是说明（合并单元格），第2行是表头

PROBLEM_HEADERS = ["比特窗口序号", "窗口名称", "FB账号名", "主页名称", "主页ID", "问题类型", "详细说明", "时间"]
PROBLEM_WIDTHS = [13, 20, 20, 24, 20, 16, 60, 18]


def next_run_number(out_dir, date_str):
    n = 1
    while (Path(out_dir) / f"{date_str}_第{n}次_FB数据.xlsx").exists():
        n += 1
    return n


def write_main(template_path, out_path, rows):
    """rows: [[达人id, 平台, 名称, 主页ID, 链接, 粉丝数], ...]"""
    wb = load_workbook(template_path)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW):
        for cell in row:
            cell.value = None
    for i, values in enumerate(rows):
        for j, v in enumerate(values):
            ws.cell(row=FIRST_DATA_ROW + i, column=1 + j, value=v)
    wb.save(out_path)


def write_problems(out_path, problems):
    """problems: [dict(seq, window, account, page_name, page_id, kind, detail, time), ...]"""
    wb = Workbook()
    ws = wb.active
    ws.title = "问题清单"
    ws.append(PROBLEM_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for i, w in enumerate(PROBLEM_WIDTHS):
        ws.column_dimensions[chr(ord("A") + i)].width = w
    for p in problems:
        ws.append([
            p.get("seq"), p.get("window", ""), p.get("account", ""), p.get("page_name", ""),
            p.get("page_id", ""), p.get("kind", ""), p.get("detail", ""), p.get("time", ""),
        ])
    for row in ws.iter_rows(min_row=2):
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(out_path)
